import os
import shutil
import subprocess
from uuid import uuid4

from fastapi import HTTPException, status
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from app.ai.ppt_generator import PPTGenerator
from app.ai.qdrant import get_paper_chunks
from app.modules.papers.repository import PaperRepository


class PPTService:
    def __init__(self, repo: PaperRepository):
        self.repo = repo
        self.generator = PPTGenerator()

    async def generate(self, paper_id, slides, current_user):
        # -----------------------------------------
        # 1. Get paper
        # -----------------------------------------
        paper = await self.repo.get_by_id(paper_id)

        if paper is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Research paper not found.",
            )

        # Prevent one user from generating a deck out of another
        # user's private paper.
        if paper.owner_id != current_user.id:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Access denied.",
            )

        # -----------------------------------------
        # 2. Check processing status
        # -----------------------------------------
        if paper.processing_status != "completed":
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="Research paper is not fully processed yet.",
            )

        # -----------------------------------------
        # 3. Retrieve indexed paper content
        # -----------------------------------------
        content = self._get_paper_content(
            paper_id
        )

        if not content:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="No indexed content found for this research paper.",
            )

        # -----------------------------------------
        # 4. Generate AI presentation outline
        # -----------------------------------------
        try:
            outline = self.generator.generate_outline(
                title=paper.title,
                abstract=paper.abstract or "",
                content=content,
                slides=slides,
            )
        except Exception as e:
            print(f"[PPTGenerator] Outline generation failed: {e}")
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail="AI failed to generate a presentation outline. Please try again.",
            ) from e

        if not outline:
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail="AI failed to generate a presentation outline.",
            )

        # -----------------------------------------
        # 5. Create PowerPoint
        # -----------------------------------------
        prs = Presentation()

        # 16:9 widescreen
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        # -----------------------------------------
        # 6. Add slides
        # -----------------------------------------
        for slide_data in outline:
            self._add_slide(
                prs,
                slide_data,
            )

        # -----------------------------------------
        # 7. Create generated directory
        # -----------------------------------------
        os.makedirs(
            "generated",
            exist_ok=True,
        )

        presentation_id = str(uuid4())

        pptx_filename = (
            f"{presentation_id}.pptx"
        )

        pptx_path = os.path.join(
            "generated",
            pptx_filename,
        )

        # -----------------------------------------
        # 8. Save PPTX
        # -----------------------------------------
        prs.save(pptx_path)

        # -----------------------------------------
        # 9. Convert PPTX → PDF
        # -----------------------------------------
        pdf_filename = (
            f"{presentation_id}.pdf"
        )

        pdf_path = os.path.join(
            "generated",
            pdf_filename,
        )

        try:
            self._convert_to_pdf(
                pptx_path,
                "generated",
            )
        except RuntimeError as e:
            print(f"[PPTGenerator] PDF conversion failed: {e}")
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=str(e),
            ) from e

        # -----------------------------------------
        # 10. Verify PDF was created
        # -----------------------------------------
        if not os.path.exists(pdf_path):
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=(
                    "PowerPoint was generated, but PDF conversion failed."
                ),
            )

        # -----------------------------------------
        # 11. Return both files
        # -----------------------------------------
        return {
            "pptx_url": (
                f"/generated/{pptx_filename}"
            ),
            "pdf_url": (
                f"/generated/{pdf_filename}"
            ),
        }

    # =========================================================
    # PDF CONVERSION
    # =========================================================

    def _convert_to_pdf(
        self,
        pptx_path,
        output_directory,
    ):
        """
        Convert PowerPoint to PDF using LibreOffice.
        """

        soffice = shutil.which(
            "soffice"
        )

        if soffice is None:

            possible_paths = [
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            ]

            for path in possible_paths:
                if os.path.exists(path):
                    soffice = path
                    break

        if soffice is None:
            raise RuntimeError(
                "LibreOffice was not found. "
                "Please make sure LibreOffice is installed."
            )

        command = [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            os.path.abspath(
                output_directory
            ),
            os.path.abspath(
                pptx_path
            ),
        ]

        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=120,
            )

        except subprocess.TimeoutExpired as exc:
            raise RuntimeError(
                "LibreOffice PDF conversion timed out."
            ) from exc

        if result.returncode != 0:
            raise RuntimeError(
                "LibreOffice failed to convert "
                f"PPTX to PDF.\n{result.stderr}"
            )

        print(
            "LibreOffice PDF conversion:",
            result.stdout.strip(),
        )

    # =========================================================
    # PAPER CONTENT
    # =========================================================

    def _get_paper_content(
        self,
        paper_id,
    ):
        """
        Retrieve all indexed chunks belonging
        to the paper in their original order.
        """

        chunks = get_paper_chunks(
            paper_id=str(paper_id)
        )

        if not chunks:
            return ""

        return "\n\n".join(
            chunk["text"]
            for chunk in chunks
            if chunk.get("text")
        )

    # =========================================================
    # SLIDE ROUTER
    # =========================================================

    def _add_slide(
        self,
        prs,
        slide_data,
    ):
        layout = slide_data.get(
            "layout",
            "content",
        )

        if layout == "title":
            self._add_title_slide(
                prs,
                slide_data,
            )

        elif layout == "two_column":
            self._add_two_column_slide(
                prs,
                slide_data,
            )

        elif layout in [
            "table",
            "comparison",
        ]:
            self._add_table_slide(
                prs,
                slide_data,
            )

        elif layout == "chart":
            self._add_chart_slide(
                prs,
                slide_data,
            )

        elif layout == "conclusion":
            self._add_conclusion_slide(
                prs,
                slide_data,
            )

        else:
            self._add_content_slide(
                prs,
                slide_data,
            )

    # =========================================================
    # COLORS
    # =========================================================

    def _color(
        self,
        hex_value,
    ):
        hex_value = hex_value.replace(
            "#",
            "",
        )

        return RGBColor(
            int(hex_value[0:2], 16),
            int(hex_value[2:4], 16),
            int(hex_value[4:6], 16),
        )

    # =========================================================
    # BACKGROUND
    # =========================================================

    def _add_background(
        self,
        slide,
    ):
        fill = slide.background.fill

        fill.solid()

        fill.fore_color.rgb = self._color(
            "FFFFFF"
        )

    # =========================================================
    # TOP ACCENT
    # =========================================================

    def _add_top_accent(
        self,
        slide,
    ):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0,
            0,
            Inches(13.333),
            Inches(0.08),
        )

        shape.fill.solid()

        shape.fill.fore_color.rgb = self._color(
            "4F46E5"
        )

        shape.line.fill.background()

    # =========================================================
    # FOOTER
    # =========================================================

    def _add_footer(
        self,
        slide,
        slide_number,
    ):
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.65),
            Inches(6.98),
            Inches(12.0),
            Inches(0.01),
        )

        line.fill.solid()

        line.fill.fore_color.rgb = self._color(
            "E5E7EB"
        )

        line.line.fill.background()

        footer = slide.shapes.add_textbox(
            Inches(0.65),
            Inches(7.08),
            Inches(5.5),
            Inches(0.2),
        )

        paragraph = (
            footer.text_frame.paragraphs[0]
        )

        paragraph.text = (
            "Resyntra • AI Research Assistant"
        )

        paragraph.font.size = Pt(8)

        paragraph.font.color.rgb = self._color(
            "6B7280"
        )

        number = slide.shapes.add_textbox(
            Inches(11.8),
            Inches(7.08),
            Inches(0.8),
            Inches(0.2),
        )

        paragraph = (
            number.text_frame.paragraphs[0]
        )

        paragraph.text = str(
            slide_number
        )

        paragraph.alignment = PP_ALIGN.RIGHT

        paragraph.font.size = Pt(8)

        paragraph.font.color.rgb = self._color(
            "6B7280"
        )

    # =========================================================
    # TITLE
    # =========================================================

    def _add_title(
        self,
        slide,
        title,
    ):
        textbox = slide.shapes.add_textbox(
            Inches(0.7),
            Inches(0.42),
            Inches(11.9),
            Inches(0.75),
        )

        frame = textbox.text_frame

        frame.clear()

        frame.word_wrap = True

        paragraph = frame.paragraphs[0]

        paragraph.text = str(title)

        paragraph.font.size = Pt(27)

        paragraph.font.bold = True

        paragraph.font.color.rgb = self._color(
            "111827"
        )

    # =========================================================
    # BULLETS
    # =========================================================

    def _add_bullets(
        self,
        slide,
        bullets,
        left=0.9,
        top=1.6,
        width=11.5,
        height=4.9,
    ):
        textbox = slide.shapes.add_textbox(
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )

        frame = textbox.text_frame

        frame.clear()

        frame.word_wrap = True

        frame.margin_left = Inches(0.08)

        frame.margin_right = Inches(0.08)

        frame.margin_top = Inches(0.05)

        for index, bullet in enumerate(
            bullets[:5]
        ):
            if index == 0:
                paragraph = frame.paragraphs[0]
            else:
                paragraph = frame.add_paragraph()

            paragraph.text = (
                f"•  {str(bullet)}"
            )

            paragraph.font.size = Pt(19)

            paragraph.font.color.rgb = self._color(
                "374151"
            )

            paragraph.space_after = Pt(15)

            paragraph.line_spacing = 1.15

        return textbox

    # =========================================================
    # CARD
    # =========================================================

    def _add_card(
        self,
        slide,
        left,
        top,
        width,
        height,
    ):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )

        card.fill.solid()

        card.fill.fore_color.rgb = self._color(
            "F9FAFB"
        )

        card.line.color.rgb = self._color(
            "E5E7EB"
        )

        return card

    # =========================================================
    # TITLE SLIDE
    # =========================================================

    def _add_title_slide(
        self,
        prs,
        slide_data,
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self._add_background(slide)

        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0,
            0,
            Inches(0.16),
            Inches(7.5),
        )

        accent.fill.solid()

        accent.fill.fore_color.rgb = self._color(
            "4F46E5"
        )

        accent.line.fill.background()

        brand = slide.shapes.add_textbox(
            Inches(0.9),
            Inches(0.75),
            Inches(6),
            Inches(0.4),
        )

        paragraph = (
            brand.text_frame.paragraphs[0]
        )

        paragraph.text = (
            "RESYNTRA  •  AI RESEARCH ASSISTANT"
        )

        paragraph.font.size = Pt(11)

        paragraph.font.bold = True

        paragraph.font.color.rgb = self._color(
            "4F46E5"
        )

        title = slide.shapes.add_textbox(
            Inches(0.9),
            Inches(1.65),
            Inches(11.2),
            Inches(2.0),
        )

        frame = title.text_frame

        frame.clear()

        frame.word_wrap = True

        frame.vertical_anchor = (
            MSO_ANCHOR.MIDDLE
        )

        paragraph = frame.paragraphs[0]

        paragraph.text = slide_data.get(
            "title",
            "Research Presentation",
        )

        paragraph.font.size = Pt(34)

        paragraph.font.bold = True

        paragraph.font.color.rgb = self._color(
            "111827"
        )

        subtitle = slide_data.get(
            "subtitle",
            "",
        )

        if subtitle:
            box = slide.shapes.add_textbox(
                Inches(0.92),
                Inches(4.0),
                Inches(10.8),
                Inches(0.9),
            )

            paragraph = (
                box.text_frame.paragraphs[0]
            )

            paragraph.text = subtitle

            paragraph.font.size = Pt(17)

            paragraph.font.color.rgb = self._color(
                "6B7280"
            )

        bottom = slide.shapes.add_textbox(
            Inches(0.92),
            Inches(6.25),
            Inches(8),
            Inches(0.4),
        )

        paragraph = (
            bottom.text_frame.paragraphs[0]
        )

        paragraph.text = (
            "Academic Research Presentation"
        )

        paragraph.font.size = Pt(10)

        paragraph.font.color.rgb = self._color(
            "9CA3AF"
        )

    # =========================================================
    # CONTENT SLIDE
    # =========================================================

    def _add_content_slide(
        self,
        prs,
        slide_data,
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self._add_background(slide)

        self._add_top_accent(slide)

        self._add_title(
            slide,
            slide_data.get(
                "title",
                "Research Overview",
            ),
        )

        self._add_bullets(
            slide,
            slide_data.get(
                "bullets",
                [],
            ),
        )

        self._add_footer(
            slide,
            len(prs.slides),
        )

    # =========================================================
    # TWO COLUMN
    # =========================================================

    def _add_two_column_slide(
        self,
        prs,
        slide_data,
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self._add_background(slide)

        self._add_top_accent(slide)

        self._add_title(
            slide,
            slide_data.get(
                "title",
                "Research Approach",
            ),
        )

        bullets = slide_data.get(
            "bullets",
            [],
        )

        midpoint = max(
            1,
            (len(bullets) + 1) // 2,
        )

        left_bullets = bullets[:midpoint]

        right_bullets = bullets[midpoint:]

        self._add_card(
            slide,
            0.7,
            1.55,
            5.8,
            4.95,
        )

        self._add_bullets(
            slide,
            left_bullets,
            left=1.0,
            top=1.9,
            width=5.2,
            height=4.15,
        )

        self._add_card(
            slide,
            6.85,
            1.55,
            5.8,
            4.95,
        )

        self._add_bullets(
            slide,
            right_bullets,
            left=7.15,
            top=1.9,
            width=5.2,
            height=4.15,
        )

        self._add_footer(
            slide,
            len(prs.slides),
        )

    # =========================================================
    # TABLE
    # =========================================================

    def _add_table_slide(
        self,
        prs,
        slide_data,
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self._add_background(slide)

        self._add_top_accent(slide)

        self._add_title(
            slide,
            slide_data.get(
                "title",
                "Comparison",
            ),
        )

        table_data = slide_data.get(
            "table"
        )

        if not table_data:
            self._add_bullets(
                slide,
                slide_data.get(
                    "bullets",
                    [],
                ),
            )

            self._add_footer(
                slide,
                len(prs.slides),
            )

            return

        headers = table_data.get(
            "headers",
            [],
        )

        rows = table_data.get(
            "rows",
            [],
        )

        if not headers:
            self._add_footer(
                slide,
                len(prs.slides),
            )

            return

        table_shape = slide.shapes.add_table(
            len(rows) + 1,
            len(headers),
            Inches(0.8),
            Inches(1.65),
            Inches(11.75),
            Inches(4.8),
        )

        table = table_shape.table

        column_width = (
            11.75 / len(headers)
        )

        for column in table.columns:
            column.width = Inches(
                column_width
            )

        for column_index, header in enumerate(
            headers
        ):
            cell = table.cell(
                0,
                column_index,
            )

            cell.text = str(header)

            cell.fill.solid()

            cell.fill.fore_color.rgb = self._color(
                "4F46E5"
            )

            for paragraph in (
                cell.text_frame.paragraphs
            ):
                paragraph.font.size = Pt(12)

                paragraph.font.bold = True

                paragraph.font.color.rgb = self._color(
                    "FFFFFF"
                )

        for row_index, row in enumerate(
            rows,
            start=1,
        ):
            for column_index in range(
                len(headers)
            ):
                value = (
                    row[column_index]
                    if column_index < len(row)
                    else ""
                )

                cell = table.cell(
                    row_index,
                    column_index,
                )

                cell.text = str(value)

                cell.fill.solid()

                cell.fill.fore_color.rgb = self._color(
                    "F9FAFB"
                )

                for paragraph in (
                    cell.text_frame.paragraphs
                ):
                    paragraph.font.size = Pt(11)

                    paragraph.font.color.rgb = self._color(
                        "374151"
                    )

        self._add_footer(
            slide,
            len(prs.slides),
        )

    # =========================================================
    # CHART
    # =========================================================

    def _add_chart_slide(
        self,
        prs,
        slide_data,
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self._add_background(slide)

        self._add_top_accent(slide)

        self._add_title(
            slide,
            slide_data.get(
                "title",
                "Research Results",
            ),
        )

        chart_data = slide_data.get(
            "chart"
        )

        if not chart_data:
            self._add_bullets(
                slide,
                slide_data.get(
                    "bullets",
                    [],
                ),
            )

            self._add_footer(
                slide,
                len(prs.slides),
            )

            return

        labels = chart_data.get(
            "labels",
            [],
        )

        values = chart_data.get(
            "values",
            [],
        )

        if not labels or not values:
            self._add_bullets(
                slide,
                slide_data.get(
                    "bullets",
                    [],
                ),
            )

            self._add_footer(
                slide,
                len(prs.slides),
            )

            return

        chart_data_obj = ChartData()

        chart_data_obj.categories = labels

        chart_data_obj.add_series(
            "Value",
            values,
        )

        chart_type = chart_data.get(
            "type",
            "bar",
        )

        if chart_type == "line":
            chart_type_enum = (
                XL_CHART_TYPE.LINE
            )

        elif chart_type == "pie":
            chart_type_enum = (
                XL_CHART_TYPE.PIE
            )

        else:
            chart_type_enum = (
                XL_CHART_TYPE.COLUMN_CLUSTERED
            )

        chart = slide.shapes.add_chart(
            chart_type_enum,
            Inches(1.0),
            Inches(1.55),
            Inches(11.3),
            Inches(4.9),
            chart_data_obj,
        ).chart

        chart.has_legend = False

        chart.has_title = False

        if chart_type != "pie":
            chart.value_axis.has_major_gridlines = True

        self._add_footer(
            slide,
            len(prs.slides),
        )

    # =========================================================
    # CONCLUSION
    # =========================================================

    def _add_conclusion_slide(
        self,
        prs,
        slide_data,
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self._add_background(slide)

        block = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.75),
            Inches(0.85),
            Inches(11.85),
            Inches(5.75),
        )

        block.fill.solid()

        block.fill.fore_color.rgb = self._color(
            "F5F3FF"
        )

        block.line.color.rgb = self._color(
            "DDD6FE"
        )

        title = slide.shapes.add_textbox(
            Inches(1.15),
            Inches(1.3),
            Inches(10.8),
            Inches(0.8),
        )

        paragraph = (
            title.text_frame.paragraphs[0]
        )

        paragraph.text = slide_data.get(
            "title",
            "Conclusion & Key Takeaways",
        )

        paragraph.font.size = Pt(29)

        paragraph.font.bold = True

        paragraph.font.color.rgb = self._color(
            "111827"
        )

        self._add_bullets(
            slide,
            slide_data.get(
                "bullets",
                [],
            ),
            left=1.15,
            top=2.25,
            width=10.6,
            height=3.6,
        )

        self._add_footer(
            slide,
            len(prs.slides),
        )